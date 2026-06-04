"""SIDCT desktop application built with PySide6."""
from __future__ import annotations

import csv
import json
import logging
import sys
from datetime import datetime, timezone
from pathlib import Path

from sidct.batch.csv_runner import _parse_row
from sidct.enums import DimensionalCatalog, Jurisdiction, ProjectProfile, Service
from sidct.materials import available_catalogs_for_material, default_catalog_for_material
from sidct.models import FittingItem, LineInput
from sidct.project import Project, create_default_project, load_project, save_project
from sidct.reports.exports import export_project_csv, export_project_xlsx, export_rows_csv
from sidct.reports.memorial_pdf import generate_pdf
from sidct.reports.project_pdf import generate_project_pdf
from sidct.system_pipe_mapping import get_calculation_ready_materials, get_mapping_notes, get_material_options
from sidct.ui.form_behavior import (
    get_calc_mode_label,
    get_calc_modes,
    get_field_meta,
    get_form_behavior,
    list_fields,
)
from sidct.ui.theme import get_light_theme_stylesheet, get_status_color
from sidct.ui.translations import (
    RESULT_CARD_KEYS,
    SERVICE_COLORS,
    SERVICE_LABELS,
    STATUS_LABELS,
    TRANSLATIONS,
)
from sidct.ui.panels.project_tree import ProjectTreePanel
from sidct.ui.panels.results_panel import ResultsPanel


SERVICES = [service.value for service in Service]
PROFILES = [profile.value for profile in ProjectProfile if profile != ProjectProfile.CUSTOM]
CATALOGS = [catalog.value for catalog in DimensionalCatalog]
JURISDICTIONS = [jurisdiction.value for jurisdiction in Jurisdiction if jurisdiction != Jurisdiction.CUSTOM]
LIQUID_FLOW_BASES = ["m3/h", "L/s", "gpm"]
GRAVITY_FLOW_BASES = ["L/s", "m3/h"]
GAS_FLOW_BASES = ["Nm3/h", "Sm3/h", "m3/h", "kg/s"]

SHORTCUTS = {
    "new": "Ctrl+N",
    "open": "Ctrl+O",
    "save": "Ctrl+S",
    "save_as": "Ctrl+Shift+S",
    "add_line": "Ctrl+L",
    "duplicate": "Ctrl+D",
    "export_pdf": "Ctrl+P",
    "export_project_pdf": "Ctrl+Shift+P",
    "import_batch": "Ctrl+I",
}



class MissingDesktopDependency(RuntimeError):
    pass


try:
    from PySide6.QtCore import QSettings, Qt
    from PySide6.QtGui import QAction, QBrush, QColor, QFont, QKeySequence, QShortcut
    from PySide6.QtWidgets import (
        QApplication,
        QComboBox,
        QDoubleSpinBox,
        QFileDialog,
        QFormLayout,
        QFrame,
        QGridLayout,
        QGroupBox,
        QHBoxLayout,
        QLabel,
        QLineEdit,
        QListWidget,
        QListWidgetItem,
        QMainWindow,
        QMessageBox,
        QPushButton,
        QScrollArea,
        QSpinBox,
        QSplitter,
        QStatusBar,
        QTableWidget,
        QTableWidgetItem,
        QTabWidget,
        QTextEdit,
        QToolBar,
        QVBoxLayout,
        QWidget,
    )
    _HAS_QT = True
except ImportError:
    _HAS_QT = False


def _ensure_qt() -> None:
    if not _HAS_QT:
        raise MissingDesktopDependency(
            "PySide6 is required for the desktop app. Install it with: pip install PySide6"
        )


def _spin(value: float, minimum: float = -1e6, maximum: float = 1e6, step: float = 1.0, decimals: int = 3):
    widget = QDoubleSpinBox()
    widget.setRange(minimum, maximum)
    widget.setDecimals(decimals)
    widget.setSingleStep(step)
    widget.setValue(value)
    return widget


if _HAS_QT:

    class CollapsibleGroupBox(QGroupBox):
        def __init__(self, title: str, parent=None):
            super().__init__(title, parent)
            self.setCheckable(True)
            self.setChecked(True)
            self.toggled.connect(self._on_toggle)

        def _on_toggle(self, checked: bool) -> None:
            layout = self.layout()
            if layout is None:
                return
            for index in range(layout.count()):
                item = layout.itemAt(index)
                widget = item.widget() if item else None
                if widget:
                    widget.setVisible(checked)
            self.setMaximumHeight(16777215 if checked else 30)


    class _Window(QMainWindow):
        def __init__(self, main_window_ref):
            super().__init__()
            self._main = main_window_ref

        def closeEvent(self, event) -> None:
            self._main.settings.setValue("window_geometry", self.saveGeometry())
            reply = QMessageBox.question(
                self,
                self._main.tr("exit_title"),
                self._main.tr("exit_message"),
                QMessageBox.Save | QMessageBox.Discard | QMessageBox.Cancel,
            )
            if reply == QMessageBox.Save:
                if self._main.save_project():
                    event.accept()
                else:
                    event.ignore()
            elif reply == QMessageBox.Discard:
                event.accept()
            else:
                event.ignore()

else:

    class CollapsibleGroupBox:  # type: ignore[no-redef]
        pass


    class _Window:  # type: ignore[no-redef]
        pass


class MainWindow:
    def __init__(self):
        _ensure_qt()
        self.settings = QSettings("SIDCT", "SIDCT Desktop")
        self.language = self.settings.value("language", "pt-BR")
        if self.language not in TRANSLATIONS:
            self.language = "pt-BR"
        self.actions: dict[str, object] = {}
        self.form_labels: dict[str, object] = {}
        self.form_rows: dict[str, tuple[object, object]] = {}
        self.field_widgets: dict[str, object] = {}
        self.group_boxes: dict[str, object] = {}
        self.result_cards: dict[str, object] = {}
        self._results_stale = False
        self._updating_form = False
        self.recent_paths = self._load_recent_paths()

        self.window = _Window(self)
        self.window.setWindowTitle(self.tr("title"))
        screen = QApplication.primaryScreen()
        if screen:
            available = screen.availableGeometry()
            width = min(1360, int(available.width() * 0.85))
            height = min(820, int(available.height() * 0.85))
        else:
            width, height = 1360, 820
        self.window.resize(width, height)
        self.window.setMinimumSize(900, 600)
        geometry = self.settings.value("window_geometry")
        if geometry:
            self.window.restoreGeometry(geometry)
        self.project: Project = create_default_project("SIDCT Project")
        self.current_path: Path | None = None
        self.current_line_id: str | None = None
        self._build_ui()
        self._seed_default_line()
        self.refresh_project_tree()

    def _load_recent_paths(self) -> list[str]:
        raw = self.settings.value("recent_projects", "[]")
        try:
            values = json.loads(str(raw))
        except Exception:
            values = []
        if not isinstance(values, list):
            return []
        return [str(path) for path in values if path and Path(str(path)).exists()][:8]

    def _save_recent_paths(self) -> None:
        self.settings.setValue("recent_projects", json.dumps(self.recent_paths[:8]))

    def _remember_recent_path(self, path: str | Path) -> None:
        text = str(Path(path))
        self.recent_paths = [item for item in self.recent_paths if item != text]
        self.recent_paths.insert(0, text)
        self.recent_paths = self.recent_paths[:8]
        self._save_recent_paths()
        self._rebuild_recent_menu()

    def tr(self, key: str) -> str:
        return TRANSLATIONS.get(self.language, TRANSLATIONS["en"]).get(key, key)

    def service_label(self, service_id: str) -> str:
        labels = SERVICE_LABELS.get(self.language, SERVICE_LABELS["en"])
        return labels.get(service_id, service_id)

    def status_label(self, status: str) -> str:
        labels = STATUS_LABELS.get(self.language, STATUS_LABELS["en"])
        return labels.get(str(status).upper(), str(status))

    def _field_label_text(self, field_id: str, required: bool = False) -> str:
        label = get_field_meta(field_id).label(self.language)
        if required:
            return f"{label} <span style='color:#B42318;font-weight:800'>*</span>"
        return label

    def _selected_service(self) -> str:
        if not hasattr(self, "service"):
            return SERVICES[0]
        data = self.service.currentData()
        return data if data else self.service.currentText()

    def _selected_mode(self) -> str:
        if not hasattr(self, "operation_mode"):
            return "calculate_new"
        data = self.operation_mode.currentData()
        return data if data else self.operation_mode.currentText()

    def show(self) -> None:
        self.window.show()

    def _build_ui(self) -> None:

        self._apply_light_theme()
        self._build_actions()

        central = QWidget()
        root = QHBoxLayout(central)
        splitter = QSplitter(Qt.Horizontal)
        root.addWidget(splitter)

        self.project_panel = ProjectTreePanel(self.language, self.tr, self.service_label, self.status_label)
        self.project_panel.lineSelected.connect(self._select_line_by_id)
        self.project_panel.removeRequested.connect(self.remove_line)
        self.project_panel.recalculateRequested.connect(self.calculate_current)
        self.project_label = self.project_panel.label
        self.project_tree = self.project_panel.tree
        splitter.addWidget(self.project_panel)

        self.form_panel = self._build_form()
        form_scroll = QScrollArea()
        form_scroll.setWidget(self.form_panel)
        form_scroll.setWidgetResizable(True)
        form_scroll.setFrameShape(QFrame.NoFrame)
        splitter.addWidget(form_scroll)

        self.results_panel = ResultsPanel(self.language, self.tr, self.status_label)
        self.result_cards = self.results_panel.result_cards
        self.tabs = self.results_panel.tabs
        self.summary_table = self.results_panel.summary_table
        self.results_table = self.results_panel.results_table
        self.warnings_box = self.results_panel.warnings_box
        self.assumptions_box = self.results_panel.assumptions_box
        self.audit_box = self.results_panel.audit_box
        splitter.addWidget(self.results_panel)
        total = self.window.width()
        splitter.setSizes([int(total * 0.19), int(total * 0.41), int(total * 0.40)])

        self.window.setCentralWidget(central)
        self.window.setStatusBar(QStatusBar())
        self._set_status(self.tr("ready"))

    def _apply_light_theme(self) -> None:
        self.window.setStyleSheet(get_light_theme_stylesheet())

    def _build_actions(self) -> None:
        toolbar = QToolBar("Main")
        self.window.addToolBar(toolbar)
        menu_file = self.window.menuBar().addMenu(self.tr("file"))
        menu_project = self.window.menuBar().addMenu(self.tr("project_menu"))
        menu_export = self.window.menuBar().addMenu(self.tr("export"))
        menu_options = self.window.menuBar().addMenu(self.tr("options"))
        menu_help = self.window.menuBar().addMenu(self.tr("about"))
        self.menus = {
            "file": menu_file,
            "project_menu": menu_project,
            "export": menu_export,
            "options": menu_options,
            "about": menu_help,
        }

        action_groups = [
            [
                ("new", self.new_project, menu_file),
                ("open", self.open_project, menu_file),
                ("save", self.save_project, menu_file),
                ("save_as", self.save_project_as, menu_file),
            ],
            [
                ("add_line", self.add_line, menu_project),
                ("duplicate", self.duplicate_line, menu_project),
                ("recalculate", self.calculate_current, menu_project),
                ("remove_line", self.remove_line, menu_project),
                ("import_batch", self.import_batch, menu_project),
            ],
            [
                ("export_pdf", self.export_pdf, menu_export),
                ("export_project_pdf", self.export_project_pdf, menu_export),
                ("export_csv", self.export_csv, menu_export),
                ("export_xlsx", self.export_xlsx, menu_export),
            ],
        ]
        for group_index, actions in enumerate(action_groups):
            if group_index:
                toolbar.addSeparator()
            for key, slot, menu in actions:
                act = QAction(self.tr(key), self.window)
                if key in SHORTCUTS:
                    act.setShortcut(SHORTCUTS[key])
                act.triggered.connect(slot)
                menu.addAction(act)
                toolbar.addAction(act)
                self.actions[key] = act

        self.recent_menu = menu_file.addMenu(self.tr("recent_projects"))
        self.menus["recent_projects"] = self.recent_menu
        self._rebuild_recent_menu()

        about_act = QAction(self.tr("about"), self.window)
        about_act.triggered.connect(self.show_about)
        menu_help.addAction(about_act)
        self.actions["about"] = about_act

        calc_shortcut = QShortcut(QKeySequence("F5"), self.window)
        calc_shortcut.activated.connect(self.calculate_current)
        self.actions["calculate_shortcut"] = calc_shortcut

        language_menu = menu_options.addMenu(self.tr("language"))
        self.menus["language"] = language_menu
        for language_key, label in [("en", "English"), ("pt-BR", "Portugues (PT-BR)")]:
            act = QAction(label, self.window)
            act.setCheckable(True)
            act.setChecked(self.language == language_key)
            act.triggered.connect(lambda checked=False, lang=language_key: self.set_language(lang))
            language_menu.addAction(act)
            self.actions[f"language_{language_key}"] = act

    def _rebuild_recent_menu(self) -> None:
        if not hasattr(self, "recent_menu"):
            return
        self.recent_menu.clear()
        for path in self.recent_paths:
            act = QAction(Path(path).name, self.window)
            act.setToolTip(path)
            act.triggered.connect(lambda checked=False, p=path: self.open_recent_project(p))
            self.recent_menu.addAction(act)
        if self.recent_paths:
            self.recent_menu.addSeparator()
        clear_act = QAction(self.tr("clear_recent"), self.window)
        clear_act.triggered.connect(self.clear_recent_projects)
        self.recent_menu.addAction(clear_act)

    def set_language(self, language: str) -> None:
        if language not in TRANSLATIONS:
            return
        self.language = language
        self.settings.setValue("language", language)
        self._apply_language()

    def _apply_language(self) -> None:
        self.window.setWindowTitle(self.tr("title"))
        for key, menu in getattr(self, "menus", {}).items():
            if key in ("language",):
                menu.setTitle(self.tr(key))
            elif key in TRANSLATIONS[self.language]:
                menu.setTitle(self.tr(key))
        for key, action in self.actions.items():
            if key.startswith("language_"):
                language = key.removeprefix("language_")
                action.setChecked(language == self.language)
            elif key in TRANSLATIONS[self.language]:
                action.setText(self.tr(key))
        if hasattr(self, "service"):
            self._refresh_service_labels()
        if hasattr(self, "operation_mode"):
            self._refresh_mode_labels()
        for key, box in self.group_boxes.items():
            box.setTitle(self.tr(key))
        if hasattr(self, "calculate_button"):
            self.calculate_button.setText(self.tr("calculate"))
        if "guidance" in self.form_labels:
            self.form_labels["guidance"].setText(self.tr("guidance"))
        if self.result_cards:
            for key, titles in RESULT_CARD_KEYS:
                title_key = f"{key}_title"
                if title_key in self.result_cards:
                    self.result_cards[title_key].setText(titles.get(self.language, titles["en"]))
        if hasattr(self, "results_panel"):
            self.results_panel.set_language(self.language)
        if hasattr(self, "tabs"):
            self.tabs.setTabText(0, self.tr("summary"))
            self.tabs.setTabText(1, self.tr("results"))
            self.tabs.setTabText(2, self.tr("warnings"))
            self.tabs.setTabText(3, self.tr("assumptions"))
            if self.tabs.count() > 4:
                self.tabs.setTabText(4, self.tr("audit"))
        if hasattr(self, "summary_table"):
            self.summary_table.setHorizontalHeaderLabels([self.tr("field"), self.tr("value")])
        if hasattr(self, "results_table"):
            self.results_table.setHorizontalHeaderLabels([self.tr("metric"), self.tr("value")])
        self._rebuild_recent_menu()
        if hasattr(self, "material"):
            self._refresh_material_options(self._selected_material())
        if hasattr(self, "field_widgets"):
            self._apply_form_behavior()
        self.refresh_project_tree()
        self._set_status(self.tr("ready"))

    def _build_form(self):

        panel = QWidget()
        layout = QVBoxLayout(panel)

        def group(title: str):
            box = CollapsibleGroupBox(title)
            form = QFormLayout(box)
            layout.addWidget(box)
            return form

        def add_row(form, field_id: str, widget) -> None:
            label = QLabel(self._field_label_text(field_id))
            label.setTextFormat(Qt.RichText)
            meta = get_field_meta(field_id)
            help_text = meta.help(self.language)
            label.setToolTip(help_text)
            widget.setToolTip(help_text)
            self.form_labels[field_id] = label
            self.form_rows[field_id] = (label, widget)
            self.field_widgets[field_id] = widget
            form.addRow(label, widget)
            self._connect_stale_signal(widget)

        self.project_name = QLineEdit("SIDCT Project")
        self.line_tag = QLineEdit("L-001")
        form = group(self.tr("project_and_line"))
        self.group_boxes["project_and_line"] = form.parentWidget()
        self.operation_mode = QComboBox()
        for mode in get_calc_modes():
            self.operation_mode.addItem(get_calc_mode_label(mode, self.language), mode)
        self.operation_mode.currentIndexChanged.connect(self._apply_form_behavior)
        add_row(form, "project_name", self.project_name)
        add_row(form, "line_tag", self.line_tag)
        add_row(form, "operation_mode", self.operation_mode)

        self.service = QComboBox()
        for service in SERVICES:
            self.service.addItem(self.service_label(service), service)
        self.service.currentIndexChanged.connect(self._service_or_region_changed)
        self.profile = QComboBox()
        self.profile.addItems(PROFILES)
        self.jurisdiction = QComboBox()
        self.jurisdiction.addItems(JURISDICTIONS)
        self.jurisdiction.currentTextChanged.connect(self._service_or_region_changed)
        self.material = QComboBox()
        self.material.currentTextChanged.connect(self._material_changed)
        self.catalog = QComboBox()
        self.catalog.addItems(CATALOGS)
        self.material_guidance = QLabel("")
        self.material_guidance.setWordWrap(True)
        form = group(self.tr("service_and_material"))
        self.group_boxes["service_and_material"] = form.parentWidget()
        add_row(form, "service", self.service)
        add_row(form, "project_profile", self.profile)
        add_row(form, "jurisdiction", self.jurisdiction)
        add_row(form, "material", self.material)
        add_row(form, "dimensional_catalog", self.catalog)
        guidance_label = QLabel(self.tr("guidance"))
        self.form_labels["guidance"] = guidance_label
        form.addRow(guidance_label, self.material_guidance)

        requirements = group("Requisitos do Sistema" if self.language == "pt-BR" else "System Requirements")
        self.group_boxes["system_requirements"] = requirements.parentWidget()
        self.requirements_panel = QLabel("")
        self.requirements_panel.setObjectName("requirementsPanel")
        self.requirements_panel.setWordWrap(True)
        requirements.addRow(self.requirements_panel)

        state_form = group(self.tr("line_state"))
        self.group_boxes["line_state"] = state_form.parentWidget()
        self.line_state_panel = QLabel("")
        self.line_state_panel.setObjectName("requirementsPanel")
        self.line_state_panel.setWordWrap(True)
        state_form.addRow(self.line_state_panel)

        form = group(self.tr("operating_conditions"))
        self.group_boxes["operating_conditions"] = form.parentWidget()
        self.P_oper = _spin(6.0, minimum=0.0, step=0.5, decimals=3)
        self.T_oper = _spin(20.0, step=1.0, decimals=2)
        self.P_design = _spin(10.0, minimum=0.0, step=0.5, decimals=3)
        self.T_design = _spin(40.0, step=1.0, decimals=2)
        add_row(form, "P_oper_bar", self.P_oper)
        add_row(form, "T_oper_c", self.T_oper)
        add_row(form, "P_design_bar", self.P_design)
        add_row(form, "T_design_c", self.T_design)

        form = group(self.tr("flow_geometry"))
        self.group_boxes["flow_geometry"] = form.parentWidget()
        self.flow = _spin(100.0, minimum=0.0001, step=1.0, decimals=3)
        self.flow_basis = QComboBox()
        self.flow_basis.addItems(LIQUID_FLOW_BASES)
        self.length = _spin(100.0, minimum=0.01, step=1.0, decimals=3)
        self.elevation = _spin(0.0, step=0.5, decimals=3)
        self.slope = _spin(10.0, minimum=0.0, step=1.0, decimals=3)
        self.vacuum = _spin(10.0, minimum=0.0, step=1.0, decimals=3)
        add_row(form, "flow_rate", self.flow)
        add_row(form, "flow_rate_basis", self.flow_basis)
        add_row(form, "line_length_m", self.length)
        add_row(form, "elevation_delta_m", self.elevation)
        add_row(form, "slope_mm_m", self.slope)
        add_row(form, "vacuum_target_mbara", self.vacuum)

        form = group(self.tr("criteria_received"))
        self.group_boxes["criteria_received"] = form.parentWidget()
        self.ca = _spin(1.5, minimum=0.0, step=0.5, decimals=3)
        self.dp_allow = _spin(0.5, minimum=0.0, step=0.05, decimals=4)
        self.dn_received = _spin(0.0, minimum=0.0, step=25.0, decimals=3)
        self.schedule_received = QLineEdit("")
        self.elbows = QSpinBox()
        self.elbows.setRange(0, 999)
        self.elbows.setValue(2)
        add_row(form, "corrosion_allowance_mm", self.ca)
        add_row(form, "allowable_pressure_drop_bar", self.dp_allow)
        add_row(form, "DN_received_mm", self.dn_received)
        add_row(form, "schedule_or_wall_received", self.schedule_received)
        add_row(form, "fittings_90lr", self.elbows)

        self.notes = QTextEdit()
        self.notes.setMaximumHeight(70)
        form = group(self.tr("notes"))
        self.group_boxes["notes"] = form.parentWidget()
        add_row(form, "design_notes", self.notes)

        self.calculate_button = QPushButton(self.tr("calculate"))
        self.calculate_button.setObjectName("calculateButton")
        self.calculate_button.clicked.connect(self.calculate_current)
        layout.addWidget(self.calculate_button)
        layout.addStretch(1)
        self._updating_form = True
        try:
            self._refresh_flow_basis_options()
            self._refresh_material_options()
            self._apply_form_behavior()
        finally:
            self._updating_form = False
        self._clear_stale()
        return panel

    def _connect_stale_signal(self, widget) -> None:
        if isinstance(widget, (QDoubleSpinBox, QSpinBox)):
            widget.valueChanged.connect(self._mark_stale)
        elif isinstance(widget, QComboBox):
            widget.currentIndexChanged.connect(self._mark_stale)
        elif isinstance(widget, QLineEdit):
            widget.textChanged.connect(self._mark_stale)
        elif isinstance(widget, QTextEdit):
            widget.textChanged.connect(self._mark_stale)

    def _mark_stale(self, *args) -> None:
        if self._updating_form or not hasattr(self, "calculate_button"):
            return
        self._results_stale = True
        self.calculate_button.setText(f"! {self.tr('calculate')}")
        self.calculate_button.setStyleSheet(
            "background: #E65100; color: white; font-weight: 700; "
            "padding: 10px 14px; border-radius: 6px;"
        )
        self._refresh_line_state()

    def _clear_stale(self) -> None:
        self._results_stale = False
        if hasattr(self, "calculate_button"):
            self.calculate_button.setText(self.tr("calculate"))
            self.calculate_button.setStyleSheet("")
        self._refresh_line_state()

    def _validation_issues(self) -> tuple[list[str], list[str], list[str]]:
        if not getattr(self, "field_widgets", None):
            return [], [], []
        behavior = self._active_form_behavior()
        pending = [
            field_id
            for field_id in behavior.required_fields
            if field_id in self.form_rows and not self._field_has_value(field_id)
        ]
        issues = []
        if self.P_design.value() < self.P_oper.value() and "P_design_bar >= P_oper_bar" in behavior.validation_rules:
            issues.append(self.tr("pressure_design_issue"))
        material_alerts = list(get_mapping_notes(self._selected_service(), self.jurisdiction.currentText())[:3])
        return pending, issues, material_alerts

    def _refresh_line_state(self) -> None:
        if not hasattr(self, "line_state_panel"):
            return
        pending, issues, material_alerts = self._validation_issues()
        pending_labels = [get_field_meta(field).label(self.language) for field in pending]
        pending_text = ", ".join(pending_labels) if pending_labels else self.tr("no_pending_fields")
        issue_text = "<br>".join(issues) if issues else self.tr("no_pending_fields")
        alert_text = "<br>".join(material_alerts) if material_alerts else self.tr("no_material_alerts")
        behavior = self._active_form_behavior()
        next_steps = [self.tr("calculate_next")]
        if self._selected_mode() == "check_received":
            next_steps.append(self.tr("review_received_next"))
        if self._current_line() and self._current_line().last_report_context:
            next_steps.append(self.tr("export_next"))
        text = (
            f"<b>{self.tr('pending_fields')}:</b> {pending_text}<br>"
            f"<b>{self.tr('applicable_criteria')}:</b> {behavior.regime(self.language)}; "
            f"{behavior.explanatory_note(self.language)}<br>"
            f"<b>{self.tr('material_alerts')}:</b> {alert_text}<br>"
            f"<b>{self.tr('next_steps')}:</b> {' | '.join(next_steps)}"
        )
        if issues:
            text += f"<br><b>{self.tr('status_field')}:</b> {issue_text}"
        self.line_state_panel.setText(text)
        self._apply_validation_styles(pending, issues)

    def _apply_validation_styles(self, pending: list[str], issues: list[str]) -> None:
        invalid = set(pending)
        if issues and "P_design_bar" in self.field_widgets:
            invalid.add("P_design_bar")
        for field_id, (label, widget) in self.form_rows.items():
            if field_id in invalid:
                widget.setStyleSheet("border: 2px solid #B42318; background: #FFF4F2;")
                label.setStyleSheet("color: #B42318; font-weight: 700;")
            else:
                widget.setStyleSheet("")
                label.setStyleSheet("")

    def _seed_default_line(self) -> None:
        route = self.project.routes[0]
        line = route.add_line(self._form_to_input())
        self.current_line_id = line.line_id

    def _refresh_service_labels(self) -> None:
        current = self._selected_service()
        self.service.blockSignals(True)
        self.service.clear()
        for service in SERVICES:
            self.service.addItem(self.service_label(service), service)
        index = self.service.findData(current)
        if index >= 0:
            self.service.setCurrentIndex(index)
        self.service.blockSignals(False)

    def _refresh_mode_labels(self) -> None:
        current = self._selected_mode()
        self.operation_mode.blockSignals(True)
        self.operation_mode.clear()
        for mode in get_calc_modes():
            self.operation_mode.addItem(get_calc_mode_label(mode, self.language), mode)
        index = self.operation_mode.findData(current)
        if index >= 0:
            self.operation_mode.setCurrentIndex(index)
        self.operation_mode.blockSignals(False)

    def _active_form_behavior(self):
        return get_form_behavior(self._selected_service(), self._selected_mode())

    def _apply_form_behavior(self) -> None:
        if not getattr(self, "field_widgets", None):
            return
        behavior = self._active_form_behavior()
        visible = set(behavior.visible_fields)
        required = set(behavior.required_fields)
        disabled = set(behavior.disabled_fields)
        warning = set(behavior.warning_fields)
        for field_id in list_fields():
            row = self.form_rows.get(field_id)
            if not row:
                continue
            label, widget = row
            is_visible = field_id in visible and field_id not in set(behavior.hidden_fields)
            label.setVisible(is_visible)
            widget.setVisible(is_visible)
            widget.setEnabled(is_visible and field_id not in disabled)
            label.setText(self._field_label_text(field_id, field_id in required))
            meta = get_field_meta(field_id)
            help_text = meta.help(self.language)
            if field_id in warning:
                help_text = f"{help_text}\n{self.tr('warning_suffix')}"
            label.setToolTip(help_text)
            widget.setToolTip(help_text)
        self._update_requirements_panel(behavior)
        self._refresh_line_state()

    def _update_requirements_panel(self, behavior) -> None:
        if not hasattr(self, "requirements_panel"):
            return
        hidden = [get_field_meta(field).label(self.language) for field in behavior.hidden_fields if field in self.form_rows]
        required = [get_field_meta(field).label(self.language) for field in behavior.required_fields if field in self.form_rows]
        if self.language == "pt-BR":
            text = (
                f"<b>Sistema:</b> {self.service_label(behavior.system_id)}<br>"
                f"<b>Regime:</b> {behavior.regime(self.language)}<br>"
                f"<b>Obrigatorios:</b> {', '.join(required) if required else 'Nenhum'}<br>"
                f"<b>Nao utilizados:</b> {', '.join(hidden[:8]) if hidden else 'Nenhum'}<br>"
                f"<b>Nota:</b> {behavior.explanatory_note(self.language)}"
            )
        else:
            text = (
                f"<b>System:</b> {self.service_label(behavior.system_id)}<br>"
                f"<b>Regime:</b> {behavior.regime(self.language)}<br>"
                f"<b>Required:</b> {', '.join(required) if required else 'None'}<br>"
                f"<b>Not used:</b> {', '.join(hidden[:8]) if hidden else 'None'}<br>"
                f"<b>Note:</b> {behavior.explanatory_note(self.language)}"
            )
        self.requirements_panel.setText(text)

    def _material_changed(self) -> None:
        material = self._selected_material()
        region = self.jurisdiction.currentText()
        catalogs = available_catalogs_for_material(material, region)
        current = self.catalog.currentText()
        self.catalog.blockSignals(True)
        self.catalog.clear()
        self.catalog.addItems(catalogs or CATALOGS)
        self.catalog.blockSignals(False)
        if current and self.catalog.findText(current) >= 0:
            self.catalog.setCurrentText(current)
            return
        catalog = default_catalog_for_material(material, region)
        if catalog:
            idx = self.catalog.findText(catalog)
            if idx >= 0:
                self.catalog.setCurrentIndex(idx)

    def _selected_material(self) -> str:
        data = self.material.currentData()
        return data if data else self.material.currentText()

    def _service_or_region_changed(self) -> None:
        self._refresh_flow_basis_options()
        self._refresh_material_options()
        self._apply_form_behavior()

    def _flow_basis_options_for_service(self, service: str) -> list[str]:
        if service in ("compressed_air", "natural_gas"):
            return GAS_FLOW_BASES
        if service in ("sanitary_drainage", "rainwater"):
            return GRAVITY_FLOW_BASES
        return LIQUID_FLOW_BASES

    def _refresh_flow_basis_options(self, preferred_basis: str | None = None) -> None:
        if not hasattr(self, "flow_basis"):
            return
        current = preferred_basis or self.flow_basis.currentText()
        options = self._flow_basis_options_for_service(self._selected_service())
        self.flow_basis.blockSignals(True)
        self.flow_basis.clear()
        self.flow_basis.addItems(options)
        idx = self.flow_basis.findText(current)
        self.flow_basis.setCurrentIndex(idx if idx >= 0 else 0)
        self.flow_basis.blockSignals(False)

    def _refresh_material_options(self, preferred_material: str | None = None) -> None:
        if not hasattr(self, "material"):
            return
        service = self._selected_service()
        region = self.jurisdiction.currentText()
        current = preferred_material or self._selected_material()
        options = get_material_options(service, region)
        ready = [option for option in options if option.calculation_ready and option.sidct_material]
        self.material.blockSignals(True)
        self.material.clear()
        for option in ready:
            label = f"{option.display_name} ({option.status})"
            self.material.addItem(label, option.sidct_material)
        if not ready:
            for material in get_calculation_ready_materials(service, region):
                self.material.addItem(material, material)
        target = -1
        for idx in range(self.material.count()):
            if self.material.itemData(idx) == current or self.material.itemText(idx) == current:
                target = idx
                break
            if self.material.itemText(idx).startswith(str(current)):
                target = idx
                break
        if target >= 0:
            self.material.setCurrentIndex(target)
        self.material.blockSignals(False)
        notes = get_mapping_notes(service, region)
        fallback = self.tr("material_guidance_empty")
        self.material_guidance.setText(" | ".join(notes[:3]) if notes else fallback)
        self._material_changed()

    def _form_to_input(self) -> LineInput:
        fittings = []
        if self.elbows.value() > 0:
            fittings.append(FittingItem(fitting_type="90_LR_ELBOW", quantity=self.elbows.value()))
        service = self._selected_service()
        mode = self._selected_mode()
        return LineInput(
            project_name=self.project_name.text().strip() or "SIDCT Project",
            line_tag=self.line_tag.text().strip() or "L-001",
            service=service,
            project_profile=self.profile.currentText(),
            jurisdiction=self.jurisdiction.currentText(),
            fluid_name=service,
            P_oper_bar=self.P_oper.value(),
            T_oper_c=self.T_oper.value(),
            P_design_bar=self.P_design.value(),
            T_design_c=self.T_design.value(),
            flow_rate=self.flow.value(),
            flow_rate_basis=self.flow_basis.currentText(),
            line_length_m=self.length.value(),
            elevation_delta_m=self.elevation.value(),
            material=self._selected_material(),
            dimensional_catalog=self.catalog.currentText(),
            corrosion_allowance_mm=self.ca.value(),
            fittings=fittings,
            allowable_pressure_drop_bar=self.dp_allow.value() or None,
            DN_received_mm=(self.dn_received.value() or None) if mode == "check_received" else None,
            schedule_or_wall_received=(
                self.schedule_received.text().strip() or None
            ) if mode == "check_received" else None,
            slope_mm_m=self.slope.value() if service in ("sanitary_drainage", "rainwater") else None,
            vacuum_target_mbara=self.vacuum.value() if service == "vacuum_utility" else None,
            design_notes=self.notes.toPlainText(),
            operation_mode=mode,
        )

    def _input_to_form(self, inp: LineInput) -> None:
        self._updating_form = True
        try:
            self.project_name.setText(inp.project_name)
            self.line_tag.setText(inp.line_tag)
            mode_idx = self.operation_mode.findData(inp.operation_mode)
            if mode_idx >= 0:
                self.operation_mode.setCurrentIndex(mode_idx)
            service_idx = self.service.findData(inp.service)
            if service_idx >= 0:
                self.service.setCurrentIndex(service_idx)
            for combo, value in [
                (self.profile, inp.project_profile),
                (self.jurisdiction, inp.jurisdiction),
                (self.catalog, inp.dimensional_catalog),
            ]:
                idx = combo.findText(value)
                if idx >= 0:
                    combo.setCurrentIndex(idx)
            self._refresh_material_options(inp.material)
            self._refresh_flow_basis_options(inp.flow_rate_basis)
            self.P_oper.setValue(inp.P_oper_bar)
            self.T_oper.setValue(inp.T_oper_c)
            self.P_design.setValue(inp.P_design_bar)
            self.T_design.setValue(inp.T_design_c)
            self.flow.setValue(inp.flow_rate)
            self.length.setValue(inp.line_length_m)
            self.elevation.setValue(inp.elevation_delta_m)
            self.ca.setValue(inp.corrosion_allowance_mm or 0.0)
            self.dp_allow.setValue(inp.allowable_pressure_drop_bar or 0.0)
            self.dn_received.setValue(inp.DN_received_mm or 0.0)
            self.schedule_received.setText(inp.schedule_or_wall_received or "")
            self.slope.setValue(inp.slope_mm_m or 0.0)
            self.vacuum.setValue(inp.vacuum_target_mbara or 0.0)
            self.notes.setPlainText(inp.design_notes)
            self._apply_form_behavior()
        finally:
            self._updating_form = False
        self._clear_stale()

    def refresh_project_tree(self) -> None:
        if hasattr(self, "project_panel"):
            self.project_panel.set_language(self.language)
            self.project_panel.refresh_tree(self.project)
            return
        self.project_tree.clear()
        self.project_label.setText(f"{self.tr('project')}: {self.project.name}")
        for route in self.project.routes:
            route_item = QListWidgetItem(f"[{self.tr('route_prefix')}] {route.name}")
            route_item.setData(256, ("route", route.route_id))
            route_font = QFont()
            route_font.setBold(True)
            route_item.setFont(route_font)
            route_item.setForeground(QBrush(QColor("#24455F")))
            self.project_tree.addItem(route_item)
            for line in route.lines:
                color = SERVICE_COLORS.get(line.line_input.service, "#94A3B8")
                status = self.status_label(line.validation_state.status)
                service = self.service_label(line.line_input.service)
                material = line.line_input.material
                item = QListWidgetItem(f"  # {line.tag}\n     {service} | {material} | {status}")
                item.setData(256, ("line", line.line_id))
                item.setForeground(QBrush(QColor("#172033")))
                item.setBackground(QBrush(QColor(color + "22")))
                item.setToolTip(f"{service} | {material} | {status}")
                self.project_tree.addItem(item)

    def _on_tree_selection(self) -> None:
        items = self.project_tree.selectedItems()
        if not items:
            return
        kind, ident = items[0].data(256)
        if kind != "line":
            return
        self._select_line_by_id(ident)

    def _select_line_by_id(self, ident: str) -> None:
        line = next((line for line in self.project.all_lines() if line.line_id == ident), None)
        if line is None:
            return
        self.current_line_id = ident
        self._input_to_form(line.line_input)
        if line.last_report_context:
            self._render_context(line.last_report_context)
        else:
            self._render_audit(line)

    def _current_line(self):
        if not getattr(self, "current_line_id", None):
            return None
        return next((line for line in self.project.all_lines() if line.line_id == self.current_line_id), None)

    def _sync_current_line(self) -> None:
        line = self._current_line()
        if line:
            line.line_input = self._form_to_input()
            line.tag = line.line_input.line_tag
            self.project.name = line.line_input.project_name

    def _field_has_value(self, field_id: str) -> bool:
        widget = self.field_widgets.get(field_id)
        if widget is None:
            return True
        if hasattr(widget, "text"):
            return bool(widget.text().strip())
        if hasattr(widget, "toPlainText"):
            return True
        if hasattr(widget, "currentText"):
            return bool(widget.currentText().strip())
        if hasattr(widget, "value"):
            value = widget.value()
            if field_id in {"flow_rate", "line_length_m", "slope_mm_m", "vacuum_target_mbara", "DN_received_mm"}:
                return value > 0
            return value is not None
        return True

    def _validate_form_before_calculation(self) -> bool:
        pending, issues, _alerts = self._validation_issues()
        missing = [get_field_meta(field_id).label(self.language) for field_id in pending]
        self._refresh_line_state()
        if missing or issues:
            lines = []
            if missing:
                lines.append(self.tr("validation_missing_intro"))
                lines.extend([f"- {name}" for name in missing])
            if issues:
                if lines:
                    lines.append("")
                lines.extend([f"- {issue}" for issue in issues])
            QMessageBox.warning(self.window, self.tr("validation_missing_title"), "\n".join(lines))
            self._set_status(self.tr("validation_missing_title"))
            return False
        return True

    def calculate_current(self) -> None:
        line = self._current_line()
        if line is None:
            return
        if not self._validate_form_before_calculation():
            return
        try:
            self._sync_current_line()
            ctx = line.calculate()
            self._stamp_context_audit(ctx)
            self.project.add_audit(
                "calculate_line",
                line.line_id,
                line.tag,
                f"Line calculated: {line.tag}",
                {"status": line.validation_state.status},
            )
            self._render_context(ctx)
            self.refresh_project_tree()
            self._auto_save()
            self._clear_stale()
            self._set_status(self.tr("calculated_line").format(tag=line.tag))
        except Exception as exc:
            self._error(self.tr("calc_failed"), str(exc))

    def _auto_save(self) -> None:
        if not self.current_path:
            return
        auto_path = self.current_path.with_suffix(".autosave.json")
        try:
            save_project(self.project, auto_path)
        except Exception:
            pass

    def _stamp_context_audit(self, ctx) -> None:
        ctx.dataset_provenance = dict(ctx.dataset_provenance or {})
        ctx.dataset_provenance["audit"] = {
            "calculated_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
            "project_updated_at": self.project.updated_at,
            "schema_version": self.project.schema_version,
            "catalog": ctx.line_input.dimensional_catalog if ctx.line_input else "",
            "service": ctx.line_input.service if ctx.line_input else "",
        }

    def _render_audit(self, line=None) -> None:
        if not hasattr(self, "audit_box"):
            return
        events = self.project.audit_log[-80:]
        if line is not None:
            line_events = [event for event in events if event.line_id in (None, line.line_id)]
        else:
            line_events = events
        lines = []
        if line is not None:
            last_case = line.hydraulic_cases[-1] if line.hydraulic_cases else None
            last = last_case.created_at if last_case else self.tr("never_calculated")
            lines.append(f"{self.tr('last_calculated')}: {last}")
            lines.append("")
        for event in line_events[-40:]:
            tag = f" [{event.line_tag}]" if event.line_tag else ""
            lines.append(f"{event.timestamp} | {event.action}{tag} | {event.message}")
        self.audit_box.setPlainText("\n".join(lines) if lines else self.tr("never_calculated"))

    def _render_context(self, ctx) -> None:
        hyd = ctx.hydraulic_result
        thick = ctx.thickness_result
        status = ctx.checker_result.overall_status if ctx.checker_result else "N/A"
        criterion = "N/A"
        if hyd and hyd.governing_criterion:
            criterion = hyd.governing_criterion
        elif ctx.checker_result and ctx.checker_result.governing_issue:
            criterion = ctx.checker_result.governing_issue
        if self.result_cards:
            self.result_cards["status"].setText(self.status_label(status))
            self.result_cards["dn"].setText(str(hyd.DN_governing_mm) if hyd else "N/A")
            self.result_cards["material"].setText(ctx.line_input.material if ctx.line_input else "N/A")
            self.result_cards["schedule"].setText(thick.selected_schedule if thick else "N/A")
            self.result_cards["criterion"].setText(str(criterion))
            bg = get_status_color(status)
            for key, _titles in RESULT_CARD_KEYS:
                frame = self.result_cards.get(f"{key}_frame")
                if frame:
                    frame.setStyleSheet(
                        "QFrame#resultCard { "
                        f"background: {bg}; border: 1px solid #D7E2EA; "
                        "border-radius: 8px; padding: 8px; }"
                    )
        self._fill_table(self.summary_table, [
            (self.tr("project_field"), ctx.project_name),
            (self.tr("line_field"), ctx.line_tag),
            (self.tr("status_field"), self.status_label(status)),
            (self.tr("governing_issue"), ctx.checker_result.governing_issue if ctx.checker_result else "N/A"),
            (self.tr("catalog_field"), ctx.line_input.dimensional_catalog if ctx.line_input else "N/A"),
        ])
        rows = []
        if hyd:
            hyd_labels = {
                "regime": "Regime",
                "dn": "DN governante [mm]" if self.language == "pt-BR" else "DN governing [mm]",
                "velocity": "Velocidade [m/s]" if self.language == "pt-BR" else "Velocity [m/s]",
                "dp": "dP total [bar]",
                "criterion": "Criterio" if self.language == "pt-BR" else "Criterion",
            }
            rows.extend([
                (hyd_labels["regime"], hyd.regime),
                (hyd_labels["dn"], hyd.DN_governing_mm),
                (hyd_labels["velocity"], hyd.velocity_ms),
                (hyd_labels["dp"], hyd.dp_total_bar),
                (hyd_labels["criterion"], hyd.governing_criterion),
            ])
        if thick:
            thick_labels = {
                "schedule": "Schedule",
                "selected": "Espessura selecionada [mm]" if self.language == "pt-BR" else "Selected wall [mm]",
                "required": "Espessura requerida [mm]" if self.language == "pt-BR" else "Required wall [mm]",
                "status": "Estado da espessura" if self.language == "pt-BR" else "Thickness status",
            }
            rows.extend([
                (thick_labels["schedule"], thick.selected_schedule),
                (thick_labels["selected"], thick.selected_wall_mm),
                (thick_labels["required"], thick.t_after_mill_tolerance_mm),
                (thick_labels["status"], thick.status),
            ])
        self._fill_table(self.results_table, rows)

        warnings = []
        assumptions = []
        for src in [
            ctx.fluid_properties,
            ctx.hydraulic_result,
            ctx.thickness_result,
            ctx.external_pressure_result,
            ctx.support_result,
            ctx.checker_result,
        ]:
            if src and hasattr(src, "warnings"):
                warnings.extend(src.warnings or [])
            if src and hasattr(src, "assumptions_used"):
                assumptions.extend(src.assumptions_used or [])
        warnings.extend([w.message for w in ctx.warnings])
        self.warnings_box.setPlainText("\n".join(warnings) or self.tr("no_warnings"))
        self.assumptions_box.setPlainText("\n".join(sorted(set(assumptions))) or self.tr("no_assumptions"))
        self._render_audit(self._current_line())

    def _fill_table(self, table, rows) -> None:
        table.setRowCount(len(rows))
        for row, (field, value) in enumerate(rows):
            table.setItem(row, 0, QTableWidgetItem(str(field)))
            if isinstance(value, float):
                value = f"{value:.5g}"
            table.setItem(row, 1, QTableWidgetItem("" if value is None else str(value)))
        table.resizeColumnsToContents()

    def new_project(self) -> None:
        self.project = create_default_project("SIDCT Project")
        self.current_path = None
        self.current_line_id = None
        self._seed_default_line()
        self.refresh_project_tree()
        self._clear_stale()
        self._set_status(self.tr("new_project_msg"))

    def open_project(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self.window, self.tr("open_project_dialog"), "", "SIDCT Project (*.sidct.json *.json)"
        )
        if not path:
            return
        self._open_project_path(Path(path))

    def _open_project_path(self, path: Path) -> None:
        try:
            self.project = load_project(path)
            self.current_path = path
            self.current_line_id = self.project.all_lines()[0].line_id if self.project.all_lines() else None
            self.refresh_project_tree()
            if self.current_line_id:
                line = self._current_line()
                if line:
                    self._input_to_form(line.line_input)
            self._clear_stale()
            self._remember_recent_path(path)
            self._set_status(self.tr("opened_msg").format(path=path))
        except Exception as exc:
            self._error(self.tr("open_failed"), str(exc))

    def open_recent_project(self, path: str) -> None:
        self._open_project_path(Path(path))

    def clear_recent_projects(self) -> None:
        self.recent_paths = []
        self._save_recent_paths()
        self._rebuild_recent_menu()

    def save_project(self) -> bool:
        if self.current_path is None:
            return self.save_project_as()
        self._sync_current_line()
        save_project(self.project, self.current_path)
        self._remember_recent_path(self.current_path)
        self._set_status(self.tr("saved_msg").format(path=self.current_path))
        return True

    def save_project_as(self) -> bool:
        path, _ = QFileDialog.getSaveFileName(
            self.window, self.tr("save_project_dialog"), "", "SIDCT Project (*.sidct.json)"
        )
        if not path:
            return False
        self.current_path = Path(path)
        return self.save_project()

    def add_line(self) -> None:
        self._sync_current_line()
        route = self.project.routes[0] if self.project.routes else self.project.add_route("Route A")
        count = len(self.project.all_lines()) + 1
        inp = self._form_to_input().model_copy(update={"line_tag": f"L-{count:03d}"})
        line = route.add_line(inp)
        self.project.add_audit("add_line", line.line_id, line.tag, f"Line added: {line.tag}")
        self.current_line_id = line.line_id
        self._input_to_form(inp)
        self.refresh_project_tree()

    def duplicate_line(self) -> None:
        line = self._current_line()
        if not line:
            return
        new_tag = f"{line.tag}-COPY"
        duplicated = self.project.duplicate_line(line.line_id, new_tag)
        self.current_line_id = duplicated.line_id
        self._input_to_form(duplicated.line_input)
        self.refresh_project_tree()

    def remove_line(self) -> None:
        line = self._current_line()
        if not line:
            return
        reply = QMessageBox.question(
            self.window,
            self.tr("confirm_remove_title"),
            self.tr("confirm_remove_msg").format(tag=line.tag),
            QMessageBox.Yes | QMessageBox.No,
        )
        if reply != QMessageBox.Yes:
            return
        removed_tag = line.tag
        self.project.remove_line(line.line_id)
        remaining = self.project.all_lines()
        self.current_line_id = remaining[0].line_id if remaining else None
        if self.current_line_id:
            self._input_to_form(remaining[0].line_input)
        self.refresh_project_tree()
        self._set_status(self.tr("removed_msg").format(tag=removed_tag))

    def import_batch(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self.window,
            self.tr("batch_dialog"),
            "",
            "Batch (*.csv *.xlsx)",
        )
        if not path:
            return
        try:
            rows = self._read_batch_rows(Path(path))
            if not rows:
                QMessageBox.information(self.window, self.tr("import_batch"), self.tr("batch_no_rows"))
                return
            route = self.project.routes[0] if self.project.routes else self.project.add_route("Route A")
            errors: list[dict[str, object]] = []
            ok = 0
            for index, row in enumerate(rows, start=1):
                line_tag = row.get("line_tag") or f"L-B{index:03d}"
                try:
                    inp = _parse_row(row, index)
                    line = route.add_line(inp)
                    self.project.add_audit("batch_add_line", line.line_id, line.tag, f"Batch line added: {line.tag}")
                    ctx = line.calculate(case_name="Batch import")
                    self._stamp_context_audit(ctx)
                    self.project.add_audit(
                        "batch_calculate_line",
                        line.line_id,
                        line.tag,
                        f"Batch line calculated: {line.tag}",
                        {"status": line.validation_state.status},
                    )
                    ok += 1
                except Exception as exc:
                    errors.append({"row_index": index, "line_tag": line_tag, "error": str(exc)})
            self.project.touch()
            if self.project.all_lines():
                self.current_line_id = self.project.all_lines()[-1].line_id
                self._input_to_form(self.project.all_lines()[-1].line_input)
            self.refresh_project_tree()
            message = self.tr("batch_done").format(ok=ok, errors=len(errors))
            if errors:
                error_path = Path(path).with_suffix(".sidct_errors.csv")
                export_rows_csv(errors, error_path)
                message = f"{message}\n{self.tr('batch_errors_saved').format(path=error_path)}"
            QMessageBox.information(self.window, self.tr("import_batch"), message)
            self._set_status(message.replace("\n", " | "))
        except Exception as exc:
            self._error(self.tr("import_batch"), str(exc))

    def _read_batch_rows(self, path: Path) -> list[dict[str, str]]:
        if path.suffix.lower() == ".csv":
            with path.open(encoding="utf-8-sig", newline="") as handle:
                return [dict(row) for row in csv.DictReader(handle)]
        if path.suffix.lower() == ".xlsx":
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise ImportError("openpyxl is required for XLSX batch import") from exc
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return []
            headers = [str(value).strip() if value is not None else "" for value in rows[0]]
            parsed: list[dict[str, str]] = []
            for raw in rows[1:]:
                item = {
                    headers[index]: "" if value is None else str(value)
                    for index, value in enumerate(raw)
                    if index < len(headers) and headers[index]
                }
                if any(str(value).strip() for value in item.values()):
                    parsed.append(item)
            return parsed
        raise ValueError(f"Unsupported batch file: {path.suffix}")

    def export_pdf(self) -> None:
        line = self._current_line()
        if not line:
            return
        if line.last_report_context is None:
            self.calculate_current()
        if line.last_report_context is None:
            return
        path, _ = QFileDialog.getSaveFileName(
            self.window, self.tr("export_pdf_dialog"), f"{line.tag}.pdf", "PDF (*.pdf)"
        )
        if path:
            generate_pdf(line.last_report_context, path)
            self._set_status(self.tr("exported_msg").format(path=path))

    def export_project_pdf(self) -> None:
        if not self.project.all_lines():
            return
        for line in self.project.all_lines():
            if line.last_report_context is None:
                try:
                    ctx = line.calculate(case_name="Project PDF export")
                    self._stamp_context_audit(ctx)
                except Exception as exc:
                    line.validation_state.errors.append(str(exc))
        self.project.touch()
        path, _ = QFileDialog.getSaveFileName(
            self.window,
            self.tr("export_project_pdf_dialog"),
            f"{self.project.name}.pdf",
            "PDF (*.pdf)",
        )
        if path:
            generate_project_pdf(self.project, path)
            self.refresh_project_tree()
            self._set_status(self.tr("exported_msg").format(path=path))

    def export_csv(self) -> None:
        path, _ = QFileDialog.getSaveFileName(
            self.window, self.tr("export_csv_dialog"), "sidct_results.csv", "CSV (*.csv)"
        )
        if path:
            export_project_csv(self.project, path)
            self._set_status(self.tr("exported_msg").format(path=path))

    def export_xlsx(self) -> None:
        path, _ = QFileDialog.getSaveFileName(
            self.window, self.tr("export_xlsx_dialog"), "sidct_results.xlsx", "Excel (*.xlsx)"
        )
        if path:
            export_project_xlsx(self.project, path)
            self._set_status(self.tr("exported_msg").format(path=path))

    def _error(self, title: str, message: str) -> None:
        QMessageBox.critical(self.window, title, message)
        self._set_status(message)

    def show_about(self) -> None:
        QMessageBox.information(self.window, self.tr("about_title"), self.tr("about_message"))

    def _set_status(self, message: str) -> None:
        lines_count = len(self.project.all_lines()) if hasattr(self, "project") else 0
        saved = "saved" if self.current_path else "unsaved"
        full = (
            f"{self.tr(saved + '_indicator')}  {message}   |   "
            f"{self.tr('lines_count')}: {lines_count}   |   {self.language.upper()}"
        )
        self.window.statusBar().showMessage(full)


def main() -> int:
    _ensure_qt()
    _activate_file_logging()
    app = QApplication.instance() or QApplication(sys.argv)
    app.setApplicationName("SIDCT")
    window = MainWindow()
    window.show()
    return app.exec()


def _activate_file_logging() -> None:
    """Configure a rotating daily log file under ~/.sidct/logs/."""
    try:
        from sidct.logging_config import setup_logging
        log_dir = Path.home() / ".sidct" / "logs"
        log_dir.mkdir(parents=True, exist_ok=True)
        date_str = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        log_file = str(log_dir / f"sidct_{date_str}.log")
        setup_logging(level="INFO", log_file=log_file)
        logging.getLogger("sidct").info(
            "SIDCT desktop started — log: %s", log_file
        )
    except Exception:  # noqa: BLE001
        pass  # logging failure must never prevent startup


if __name__ == "__main__":
    raise SystemExit(main())
