import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from sidct.batch.csv_runner import run_batch_from_bytes
from sidct.catalogs.pipe_dimension_catalog import get_available_dns, get_pipe_dimension
from sidct.engines.selector import run_full_calculation
from sidct.materials import validate_material_application
from sidct.models import LineInput
from sidct.project import create_default_project, load_project, save_project
from sidct.reports.exports import export_project_csv, export_project_xlsx
from sidct.reports.memorial_pdf import generate_pdf
from sidct.reports.project_pdf import generate_project_pdf


def _line(tag: str, material: str, catalog: str, jurisdiction: str = "EU") -> LineInput:
    return LineInput(
        project_name="PROD",
        line_tag=tag,
        service="service_water",
        project_profile="industrial_utilities_eu",
        jurisdiction=jurisdiction,
        fluid_name="water",
        P_oper_bar=4.0,
        T_oper_c=20.0,
        P_design_bar=8.0,
        T_design_c=30.0,
        flow_rate=25.0,
        flow_rate_basis="m3/h",
        line_length_m=80.0,
        material=material,
        dimensional_catalog=catalog,
        corrosion_allowance_mm=0.0 if "PVC" in material.upper() else 1.5,
    )


def test_project_persistence_two_independent_routes(tmp_path):
    project = create_default_project("Plant A")
    route_a = project.routes[0]
    route_b = project.add_route("Route B")
    line_a = route_a.add_line(_line("CS-001", "A106 GrB", "ASME_B36_10M"))
    line_b = route_b.add_line(_line("SS-001", "A312 TP316", "ASME_B36_19M"))

    project.calculate_line(line_a.line_id)
    project.calculate_line(line_b.line_id)

    path = tmp_path / "plant.sidct.json"
    save_project(project, path)
    loaded = load_project(path)

    assert len(loaded.routes) == 2
    assert loaded.routes[0].lines[0].tag == "CS-001"
    assert loaded.routes[1].lines[0].tag == "SS-001"
    assert loaded.routes[0].lines[0].line_input.material == "A106 GrB"
    assert loaded.routes[1].lines[0].line_input.material == "A312 TP316"


def test_pvc_catalogs_eu_and_us_are_distinct():
    eu = get_pipe_dimension("PVC_EN1452", 100, "PN16")
    us = get_pipe_dimension("PVC_ASTMD1785", 100, "SCH40")

    assert eu.catalog == "PVC_EN1452"
    assert us.catalog == "PVC_ASTMD1785"
    assert eu.OD_mm != us.OD_mm
    assert 100 in get_available_dns("PVC_EN1452")
    assert 100 in get_available_dns("PVC_ASTMD1785")


def test_pvc_water_calculation_uses_pvc_catalog():
    inp = _line("PVC-001", "PVC", "PVC_EN1452")
    ctx = run_full_calculation(inp)

    assert ctx.line_input.dimensional_catalog == "PVC_EN1452"
    assert ctx.hydraulic_result is not None
    assert ctx.hydraulic_result.DN_governing_mm is not None
    assert ctx.thickness_result is not None
    assert ctx.thickness_result.status == "CALCULATED"


def test_pe100_water_calculation_uses_en12201_catalog():
    inp = _line("PE-001", "HDPE", "ASME_B36_10M")
    ctx = run_full_calculation(inp)

    assert ctx.line_input.dimensional_catalog == "PE_EN12201"
    assert ctx.hydraulic_result is not None
    assert ctx.hydraulic_result.DN_governing_mm is not None
    assert ctx.thickness_result is not None
    assert ctx.thickness_result.status == "CALCULATED"
    assert (ctx.thickness_result.selected_schedule or "").startswith("SDR")


def test_ppr_water_calculation_uses_iso15874_catalog():
    inp = _line("PPR-001", "PPR", "ASME_B36_10M")
    ctx = run_full_calculation(inp)

    assert ctx.line_input.dimensional_catalog == "PPR_ISO15874"
    assert ctx.hydraulic_result is not None
    assert ctx.hydraulic_result.DN_governing_mm is not None
    assert ctx.thickness_result is not None
    assert ctx.thickness_result.status == "CALCULATED"
    assert (ctx.thickness_result.selected_schedule or "").startswith("SDR")


def test_pvc_us_profile_uses_us_catalog_when_steel_catalog_was_requested():
    inp = _line("PVC-US-001", "PVC", "ASME_B36_10M", jurisdiction="US")
    ctx = run_full_calculation(inp)

    assert ctx.line_input.dimensional_catalog == "PVC_ASTMD1785"
    assert any("Using 'PVC_ASTMD1785'" in w.message for w in ctx.warnings)


def test_material_inadequate_for_compressed_air_is_blocked():
    with pytest.raises(Exception):
        validate_material_application("PVC", "compressed_air", 30.0, 8.0, "EU")


def test_project_exports_and_pdf(tmp_path):
    project = create_default_project("Exports")
    line = project.routes[0].add_line(_line("CS-EXP", "A106 GrB", "ASME_B36_10M"))
    ctx = project.calculate_line(line.line_id)

    csv_path = tmp_path / "results.csv"
    xlsx_path = tmp_path / "results.xlsx"
    pdf_path = tmp_path / "line.pdf"
    export_project_csv(project, csv_path)
    export_project_xlsx(project, xlsx_path)
    generate_pdf(ctx, pdf_path)

    assert csv_path.exists() and csv_path.read_text(encoding="utf-8").startswith("project_name")
    assert xlsx_path.exists() and xlsx_path.stat().st_size > 0
    assert pdf_path.exists() and pdf_path.stat().st_size > 0

    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path)
    assert {"SIDCT Results", "Line Inputs", "Audit", "Warnings"}.issubset(set(wb.sheetnames))
    result_sheet = wb["SIDCT Results"]
    headers = [cell.value for cell in result_sheet[1]]
    values = [cell.value for cell in result_sheet[2]]
    assert headers[:3] == ["Project", "Route", "Line Tag"]
    assert headers[headers.index("Velocity [m/s]")] == "Velocity [m/s]"
    assert isinstance(values[headers.index("Velocity [m/s]")], str)
    assert len(values[headers.index("Velocity [m/s]")].split(".")[-1]) <= 3
    assert result_sheet.freeze_panes == "A2"


def test_gravity_exports_include_drainage_specific_results(tmp_path):
    project = create_default_project("Drainage Export")
    inp = LineInput(
        project_name="Drainage Export",
        line_tag="DR-001",
        service="sanitary_drainage",
        project_profile="glass_factory_industrial_eu",
        jurisdiction="EU",
        fluid_name="sanitary_drainage",
        P_oper_bar=0.0,
        T_oper_c=20.0,
        P_design_bar=0.0,
        T_design_c=40.0,
        flow_rate=0.5,
        flow_rate_basis="L/s",
        line_length_m=10.0,
        material="PVC",
        dimensional_catalog="PVC_EN1452",
        corrosion_allowance_mm=0.0,
        slope_mm_m=15.0,
    )
    line = project.routes[0].add_line(inp)
    ctx = project.calculate_line(line.line_id)
    assert ctx.hydraulic_result.flow_depth_ratio is not None

    csv_path = tmp_path / "drainage.csv"
    xlsx_path = tmp_path / "drainage.xlsx"
    pdf_path = tmp_path / "drainage.pdf"
    export_project_csv(project, csv_path)
    export_project_xlsx(project, xlsx_path)
    pdf = generate_pdf(ctx, pdf_path)

    csv_text = csv_path.read_text(encoding="utf-8")
    assert "flow_depth_ratio" in csv_text
    assert "slope_adequacy" in csv_text
    assert ",N/A," in csv_text

    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, read_only=True)
    headers = [cell.value for cell in wb["SIDCT Results"][1]]
    assert "Flow Depth Ratio" in headers
    assert "Slope Adequacy" in headers
    assert "Self Cleansing OK" in headers
    assert len(pdf) > 1000


def test_project_multi_line_pdf(tmp_path):
    project = create_default_project("Project Report")
    line_a = project.routes[0].add_line(_line("CS-PDF", "A106 GrB", "ASME_B36_10M"))
    line_b = project.routes[0].add_line(_line("PE-PDF", "HDPE", "ASME_B36_10M"))
    project.calculate_line(line_a.line_id)
    project.calculate_line(line_b.line_id)

    pdf_path = tmp_path / "project.pdf"
    pdf = generate_project_pdf(project, pdf_path)

    assert len(pdf) > 1000
    assert pdf_path.exists() and pdf_path.stat().st_size == len(pdf)


def test_batch_import_still_runs():
    csv_bytes = (
        b"project_name,line_tag,service,project_profile,jurisdiction,fluid_name,"
        b"P_oper_bar,T_oper_c,P_design_bar,T_design_c,flow_rate,flow_rate_basis,"
        b"line_length_m,material,dimensional_catalog\n"
        b"PROD,B-001,service_water,industrial_utilities_eu,EU,water,"
        b"4,20,8,30,25,m3/h,80,A106 GrB,ASME_B36_10M\n"
    )
    result_csv, errors_csv, summary = run_batch_from_bytes(csv_bytes)

    assert b"B-001" in result_csv
    assert errors_csv is None
    assert "1 linhas" in summary
