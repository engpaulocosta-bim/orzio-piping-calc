"""Tabular exports for SIDCT projects and calculation contexts."""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Iterable

from ..models import ReportContext
from ..project import Project, LineSegment, Route


RESULT_FIELDS = [
    "project_name",
    "route_name",
    "line_tag",
    "service",
    "material",
    "catalog",
    "DN_governing_mm",
    "governing_criterion",
    "velocity_ms",
    "dp_total_bar",
    "head_loss_m",
    "flow_depth_ratio",
    "slope_adequacy",
    "self_cleansing_ok",
    "selected_schedule",
    "selected_wall_mm",
    "checker_status",
    "governing_issue",
    "warnings",
    "dataset_missing",
    "out_of_scope",
    "calculated_at",
    "schema_version",
]

RESULT_LABELS = {
    "project_name": "Project",
    "route_name": "Route",
    "line_tag": "Line Tag",
    "service": "Service",
    "material": "Material",
    "catalog": "Catalog",
    "DN_governing_mm": "DN Governing [mm]",
    "governing_criterion": "Governing Criterion",
    "velocity_ms": "Velocity [m/s]",
    "dp_total_bar": "Pressure Drop [bar]",
    "head_loss_m": "Head Loss [m]",
    "flow_depth_ratio": "Flow Depth Ratio",
    "slope_adequacy": "Slope Adequacy",
    "self_cleansing_ok": "Self Cleansing OK",
    "selected_schedule": "Selected Schedule",
    "selected_wall_mm": "Selected Wall [mm]",
    "checker_status": "Checker Status",
    "governing_issue": "Governing Issue",
    "warnings": "Warnings",
    "dataset_missing": "Dataset Missing",
    "out_of_scope": "Out of Scope",
    "calculated_at": "Calculated At",
    "schema_version": "Schema Version",
}

_DECIMALS = {
    "DN_governing_mm": 0,
    "velocity_ms": 3,
    "dp_total_bar": 4,
    "head_loss_m": 3,
    "flow_depth_ratio": 3,
    "selected_wall_mm": 2,
}


def _join(items: list[str] | None) -> str:
    return "; ".join(items or [])


def _context_warnings(ctx: ReportContext) -> str:
    warnings: list[str] = [w.message for w in ctx.warnings]
    for source in [
        ctx.fluid_properties,
        ctx.hydraulic_result,
        ctx.thickness_result,
        ctx.external_pressure_result,
        ctx.support_result,
        ctx.checker_result,
    ]:
        warnings.extend(getattr(source, "warnings", []) if source else [])

    unique: list[str] = []
    for warning in warnings:
        if warning and warning not in unique:
            unique.append(warning)
    return _join(unique)


def _format_number(value: float, decimals: int) -> str:
    if decimals == 0:
        return f"{value:.0f}"
    return f"{value:.{decimals}f}".rstrip("0").rstrip(".")


def format_export_value(field: str, value: object) -> object:
    """Format values for human-readable CSV/XLSX/PDF exports."""
    if value is None or value == "":
        return "N/A"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        return _format_number(value, _DECIMALS.get(field, 4))
    if isinstance(value, list):
        return _join([str(item) for item in value])
    return value


def format_export_row(row: dict[str, object], fields: list[str] | None = None) -> dict[str, object]:
    selected_fields = fields or list(row.keys())
    return {field: format_export_value(field, row.get(field)) for field in selected_fields}


def _ordered_row(row: dict[str, object]) -> dict[str, object]:
    return {field: row.get(field) for field in RESULT_FIELDS}


def context_to_row(ctx: ReportContext) -> dict[str, object]:
    inp = ctx.line_input
    hyd = ctx.hydraulic_result
    thick = ctx.thickness_result
    check = ctx.checker_result
    audit = (ctx.dataset_provenance or {}).get("audit", {})
    return _ordered_row({
        "project_name": ctx.project_name,
        "route_name": "",
        "line_tag": ctx.line_tag,
        "service": inp.service if inp else "",
        "material": inp.material if inp else "",
        "catalog": inp.dimensional_catalog if inp else "",
        "DN_governing_mm": hyd.DN_governing_mm if hyd else None,
        "governing_criterion": hyd.governing_criterion if hyd else None,
        "velocity_ms": hyd.velocity_ms if hyd else None,
        "dp_total_bar": hyd.dp_total_bar if hyd else None,
        "head_loss_m": hyd.head_loss_m if hyd else None,
        "flow_depth_ratio": hyd.flow_depth_ratio if hyd else None,
        "slope_adequacy": hyd.slope_adequacy if hyd else None,
        "self_cleansing_ok": hyd.self_cleansing_ok if hyd else None,
        "selected_schedule": thick.selected_schedule if thick else None,
        "selected_wall_mm": thick.selected_wall_mm if thick else None,
        "checker_status": check.overall_status if check else "",
        "governing_issue": check.governing_issue if check else "",
        "warnings": _context_warnings(ctx),
        "dataset_missing": _join(check.dataset_missing_items if check else []),
        "out_of_scope": _join(check.out_of_scope_items if check else []),
        "calculated_at": audit.get("calculated_at", ""),
        "schema_version": audit.get("schema_version", ""),
    })


def _uncalculated_row(project: Project, route: Route, line: LineSegment) -> dict[str, object]:
    return _ordered_row({
        "project_name": project.name,
        "route_name": route.name,
        "line_tag": line.tag,
        "service": line.line_input.service,
        "material": line.line_input.material,
        "catalog": line.line_input.dimensional_catalog,
        "checker_status": line.validation_state.status,
        "warnings": _join(line.validation_state.warnings),
        "dataset_missing": _join(line.validation_state.dataset_missing),
        "out_of_scope": _join(line.validation_state.out_of_scope),
        "schema_version": project.schema_version,
    })


def project_rows(project: Project) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for route in project.routes:
        for line in route.lines:
            if line.last_report_context is None:
                rows.append(_uncalculated_row(project, route, line))
            else:
                row = context_to_row(line.last_report_context)
                row["route_name"] = route.name
                rows.append(_ordered_row(row))
    return rows


def export_rows_csv(rows: Iterable[dict[str, object]], path: str | Path) -> None:
    rows = list(rows)
    fieldnames = list(rows[0].keys()) if rows else ["project_name", "line_tag", "checker_status"]
    with Path(path).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(format_export_row(row, fieldnames) for row in rows)


def export_project_csv(project: Project, path: str | Path) -> None:
    export_rows_csv(project_rows(project), path)


def export_project_xlsx(project: Project, path: str | Path) -> None:
    rows = project_rows(project)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ImportError("openpyxl is required for XLSX export") from exc

    def style_sheet(ws, widths_max: int = 34) -> None:
        header_fill = PatternFill("solid", fgColor="24455F")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column_cells in ws.columns:
            letter = get_column_letter(column_cells[0].column)
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            ws.column_dimensions[letter].width = min(max(max_len + 2, 12), widths_max)

    wb = Workbook()
    ws = wb.active
    ws.title = "SIDCT Results"
    fieldnames = RESULT_FIELDS if rows else ["project_name", "line_tag", "checker_status"]
    ws.append([RESULT_LABELS.get(field, field) for field in fieldnames])
    for row in rows:
        formatted = format_export_row(row, fieldnames)
        ws.append([formatted.get(field) for field in fieldnames])
    style_sheet(ws)

    ws_inputs = wb.create_sheet("Line Inputs")
    input_fields = [
        "route_name", "project_name", "line_tag", "service", "project_profile", "jurisdiction",
        "fluid_name", "P_oper_bar", "T_oper_c", "P_design_bar", "T_design_c",
        "flow_rate", "flow_rate_basis", "line_length_m", "elevation_delta_m",
        "material", "dimensional_catalog", "design_code", "corrosion_allowance_mm",
        "allowable_pressure_drop_bar", "DN_received_mm", "schedule_or_wall_received",
        "slope_mm_m", "vacuum_target_mbara", "operation_mode", "design_notes",
    ]
    ws_inputs.append([RESULT_LABELS.get(field, field.replace("_", " ").title()) for field in input_fields])
    for route in project.routes:
        for line in route.lines:
            inp = line.line_input
            values = {
                "route_name": route.name,
                **inp.model_dump(),
            }
            ws_inputs.append([format_export_value(field, values.get(field)) for field in input_fields])
    style_sheet(ws_inputs)

    ws_audit = wb.create_sheet("Audit")
    audit_fields = ["timestamp", "action", "line_tag", "message", "details"]
    ws_audit.append([field.replace("_", " ").title() for field in audit_fields])
    for event in project.audit_log:
        ws_audit.append([
            event.timestamp,
            event.action,
            event.line_tag or "",
            event.message,
            json.dumps(event.details or {}, ensure_ascii=False, sort_keys=True),
        ])
    style_sheet(ws_audit, widths_max=50)

    ws_warnings = wb.create_sheet("Warnings")
    ws_warnings.append(["Route", "Line Tag", "Status", "Warnings", "Dataset Missing", "Out Of Scope"])
    for route in project.routes:
        for line in route.lines:
            if any([line.validation_state.warnings, line.validation_state.dataset_missing, line.validation_state.out_of_scope]):
                ws_warnings.append([
                    route.name,
                    line.tag,
                    line.validation_state.status,
                    _join(line.validation_state.warnings),
                    _join(line.validation_state.dataset_missing),
                    _join(line.validation_state.out_of_scope),
                ])
    if ws_warnings.max_row == 1:
        ws_warnings.append(["N/A", "N/A", "N/A", "No warnings or dataset gaps recorded.", "N/A", "N/A"])
    style_sheet(ws_warnings, widths_max=60)

    wb.save(path)
